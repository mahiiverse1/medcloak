"""Generate clean downloadable copies from already-redacted synthetic text."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas


DOWNLOAD_FORMATS = {"txt", "docx", "pdf"}


def build_protected_document(redacted_text: str, output_format: str) -> tuple[bytes, str, str]:
    """Return bytes, MIME type, and filename for a generated protected copy."""
    if output_format not in DOWNLOAD_FORMATS:
        raise ValueError("Choose txt, docx, or pdf.")
    if output_format == "txt":
        return redacted_text.encode("utf-8"), "text/plain; charset=utf-8", "medcloak-protected-note.txt"
    if output_format == "docx":
        return _build_docx(redacted_text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "medcloak-protected-note.docx"
    return _build_pdf(redacted_text), "application/pdf", "medcloak-protected-note.pdf"


def _build_docx(redacted_text: str) -> bytes:
    document = Document()
    document.core_properties.author = "MedCloak"
    document.core_properties.title = "MedCloak Protected Clinical Note"
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10.5)
    normal_style.paragraph_format.space_after = Pt(7)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("MedCloak Protected Clinical Note")
    title_run.font.name = "Aptos Display"
    title_run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = document.add_paragraph("De identified research prototype output")
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.runs[0]
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(9.5)
    subtitle_run.font.color.rgb = RGBColor(68, 68, 68)

    notice = document.add_paragraph()
    notice.paragraph_format.space_after = Pt(15)
    notice_run = notice.add_run("Human review is required before use or sharing. This is not a compliance determination.")
    notice_run.italic = True
    notice_run.font.size = Pt(9.5)

    heading = document.add_heading("Protected Clinical Note", level=1)
    heading.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    for paragraph_text in redacted_text.splitlines() or [redacted_text]:
        paragraph = document.add_paragraph(paragraph_text)
        paragraph.paragraph_format.space_after = Pt(8)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(10.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("MedCloak protected copy - human review required")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(90, 90, 90)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pdf(redacted_text: str) -> bytes:
    buffer = BytesIO()
    page = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    left, right, top, bottom = 56, width - 56, height - 58, 52
    page_number = 1

    def draw_page_chrome(number: int) -> float:
        page.setFillColorRGB(0, 0, 0)
        page.setFont("Helvetica-Bold", 16)
        page.drawString(left, top, "MedCloak Protected Clinical Note")
        page.setFont("Helvetica", 9)
        page.setFillColorRGB(0.28, 0.28, 0.28)
        page.drawString(left, top - 17, "De identified research prototype output")
        page.setFont("Helvetica-Oblique", 8.5)
        page.drawString(left, top - 33, "Human review is required before use or sharing. Not a compliance determination.")
        page.setFont("Helvetica", 8)
        page.setFillColorRGB(0.38, 0.38, 0.38)
        page.drawCentredString(width / 2, 30, f"MedCloak protected copy | Page {number}")
        return top - 63

    y = draw_page_chrome(page_number)
    page.setTitle("MedCloak Protected Clinical Note")
    page.setAuthor("MedCloak")
    page.setFillColorRGB(0, 0, 0)
    page.setFont("Helvetica-Bold", 11)
    page.drawString(left, y, "Protected Clinical Note")
    y -= 20
    page.setFont("Helvetica", 10)
    for source_line in redacted_text.splitlines() or [redacted_text]:
        for line in simpleSplit(source_line, "Helvetica", 10, right - left) or [""]:
            if y < bottom + 14:
                page.showPage()
                page_number += 1
                y = draw_page_chrome(page_number)
                page.setFont("Helvetica", 10)
                page.setFillColorRGB(0, 0, 0)
            page.drawString(left, y, line)
            y -= 14
        y -= 4
    page.save()
    return buffer.getvalue()
